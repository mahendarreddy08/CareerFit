import io
import os
import zipfile

from docx import Document
from pypdf import PdfReader

MAX_FILE_SIZE = 10 * 1024 * 1024


class ExtractionError(ValueError):
    """Raised when a resume file cannot be read or is not supported."""


def read_file(filename):
    with open(filename, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def _is_supported_pdf(file_bytes):
    return file_bytes.startswith(b"%PDF")


def _is_supported_docx(file_bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            return "word/document.xml" in archive.namelist()
    except zipfile.BadZipFile:
        return False


def _validate_uploaded_file(file_bytes, filename):
    if not file_bytes:
        raise ExtractionError("The uploaded resume file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ExtractionError("The uploaded resume file is too large. Please use a file smaller than 10 MB.")

    normalized_name = (filename or "").lower()

    if normalized_name.endswith(".pdf"):
        if not _is_supported_pdf(file_bytes):
            raise ExtractionError("The uploaded PDF file is unreadable or corrupt.")
        return

    if normalized_name.endswith(".docx"):
        if not _is_supported_docx(file_bytes):
            raise ExtractionError("The uploaded DOCX file is unreadable or corrupt.")
        return

    raise ExtractionError("Unsupported file type. Please upload a PDF or DOCX resume.")


def extract_text_from_pdf(file_bytes):
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text)
        return "\n".join(pages).strip()
    except Exception as exc:  # pragma: no cover - defensive branch
        raise ExtractionError("The uploaded PDF could not be read. Please upload a valid PDF resume.") from exc


def extract_text_from_docx(file_bytes):
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:  # pragma: no cover - defensive branch
        raise ExtractionError("The uploaded DOCX could not be read. Please upload a valid DOCX resume.") from exc

    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                paragraphs.append(" ".join(cells))

    return "\n".join(paragraphs).strip()


def extract_text_from_file(file_obj, filename=None):
    if hasattr(file_obj, "read"):
        file_bytes = file_obj.read()
    else:
        file_bytes = file_obj

    if isinstance(file_bytes, str):
        file_bytes = file_bytes.encode("utf-8")

    if file_bytes is None:
        raise ExtractionError("The uploaded resume file is empty.")

    resolved_filename = filename or getattr(file_obj, "filename", "") or ""
    _validate_uploaded_file(file_bytes, resolved_filename)

    normalized_name = resolved_filename.lower()
    if normalized_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if normalized_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)

    raise ExtractionError("Unsupported file type. Please upload a PDF or DOCX resume.")

